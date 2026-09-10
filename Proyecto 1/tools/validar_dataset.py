#!/usr/bin/env python3
"""Audita conteos, llaves y reglas cruzadas del libro oficial."""

from collections import Counter, defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

BOOK = Path(__file__).resolve().parents[1] / "data" / "dataset_comercial_la_estrella.xlsx"


def main():
    wb = load_workbook(BOOK, data_only=True, read_only=True)
    data = {}
    total = 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))[1:]
        data[ws.title] = rows
        total += len(rows)
        print(f"{ws.title:24} {len(rows):5}")
    assert len(wb.sheetnames) == 19 and total == 9068

    people = {row[0]: row for row in data["PERSONA"]}
    employees = {row[0]: row for row in data["EMPLEADO"]}
    sales = {row[0]: row for row in data["VENTA"]}
    states = {row[0]: row[1] for row in data["ESTADO_VENTA"]}
    catalog = {(row[2], row[3]) for row in data["CATALOGO_PRODUCTO"]}
    details = {(row[3], row[4]) for row in data["DESGLOSE_VENTA"]}
    detail_totals = defaultdict(Decimal)
    payment_totals = defaultdict(Decimal)
    for qty, price, subtotal, sale_id, _ in data["DESGLOSE_VENTA"]:
        assert Decimal(str(subtotal)) == Decimal(str(qty)) * Decimal(str(price))
        assert (sales[sale_id][2], _) in catalog
        detail_totals[sale_id] += Decimal(str(subtotal))
    for _, amount, _, sale_id in data["PAGO"]:
        payment_totals[sale_id] += Decimal(str(amount))

    assert len(catalog) == len(data["CATALOGO_PRODUCTO"])
    assert len(details) == len(data["DESGLOSE_VENTA"])
    assert all(employees[s[3]][2] == s[2] for s in data["VENTA"])
    assert all(s[0] in detail_totals for s in data["VENTA"])
    assert all(states[s[5]] != "PAGADA" or detail_totals[s[0]] == payment_totals[s[0]] for s in data["VENTA"])
    assert all(e[1].date() <= date.today() for e in data["EMPLEADO"])
    assert len([p for p in people.values() if p[4] is None]) == 45
    print(f"TOTAL{'':19} {total:5}")
    print("Todas las validaciones finalizaron correctamente.")


if __name__ == "__main__":
    main()
