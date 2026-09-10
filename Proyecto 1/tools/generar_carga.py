#!/usr/bin/env python3
"""Genera el DML Oracle a partir del libro oficial sin alterar los datos."""

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "data" / "dataset_comercial_la_estrella.xlsx"
OUTPUT = ROOT / "entrega" / "03_carga_datos.sql"

TABLES = {
    "PAIS": ("pais", ("id_pais", "nombre")),
    "DEPARTAMENTO": ("departamento", ("id_departamento", "nombre", "id_pais")),
    "MUNICIPIO": ("municipio", ("id_municipio", "nombre", "id_departamento")),
    "TIPO_TIENDA": ("tipo_tienda", ("id_tipo_tienda", "nombre")),
    "TIPO_IDENTIFICACION": ("tipo_identificacion", ("id_tipo_identificacion", "nombre")),
    "CARGO": ("cargo", ("id_cargo", "nombre")),
    "CATEGORIA": ("categoria", ("id_categoria", "nombre")),
    "MARCA": ("marca", ("id_marca", "nombre")),
    "ESTADO_VENTA": ("estado_venta", ("id_estado_venta", "nombre")),
    "METODO_PAGO": ("metodo_pago", ("id_metodo_pago", "nombre")),
    "PERSONA": ("persona", ("id_persona", "nombres", "apellidos", "telefono", "correo", "direccion", "id_municipio")),
    "TIENDA": ("tienda", ("id_tienda", "nombre", "direccion", "telefono", "id_municipio", "id_tipo_tienda")),
    "EMPLEADO": ("empleado", ("id_empleado", "fecha_contratacion", "id_tienda", "id_cargo", "id_persona")),
    "CLIENTE": ("cliente", ("id_cliente", "id_tipo_identificacion", "numero_identificacion", "id_persona")),
    "PRODUCTO": ("producto", ("id_producto", "nombre", "descripcion", "id_categoria", "id_marca")),
    "CATALOGO_PRODUCTO": ("catalogo_producto", ("precio_vigente", "existencia_actual", "id_tienda", "id_producto")),
    "VENTA": ("venta", ("id_venta", "fecha_venta", "id_tienda", "id_empleado", "id_cliente", "id_estado_venta")),
    "DESGLOSE_VENTA": ("detalle_venta", ("cantidad", "precio_unitario", "subtotal", "id_venta", "id_producto")),
    "PAGO": ("pago", ("id_pago", "monto", "id_metodo_pago", "id_venta")),
}


def oracle_literal(value):
    if value is None:
        return "NULL"
    if isinstance(value, datetime):
        return f"TO_DATE('{value:%Y-%m-%d}', 'YYYY-MM-DD')"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float, Decimal)):
        return format(Decimal(str(value)), "f")
    return "'" + str(value).replace("'", "''") + "'"


def main():
    workbook = load_workbook(INPUT, data_only=True, read_only=True)
    lines = [
        "-- Proyecto 1 - Comercial La Estrella",
        "-- Carga generada desde dataset_comercial_la_estrella.xlsx.",
        "-- Ejecutar despues de 03_creacion_DB.sql con Run Script (F5).",
        "SET DEFINE OFF;",
        "WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK;",
        "",
    ]
    total = 0
    for sheet_name, (table_name, columns) in TABLES.items():
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        count = 0
        lines.append(f"-- {sheet_name} -> {table_name}")
        for row in rows:
            values = ", ".join(oracle_literal(value) for value in row)
            lines.append(f"INSERT INTO {table_name} ({', '.join(columns)}) VALUES ({values});")
            count += 1
        total += count
        lines.extend((f"-- Filas: {count}", ""))

    lines.extend([
        "COMMIT;",
        "",
        "PROMPT Conteos posteriores a la carga",
        "SELECT tabla, filas FROM (",
    ])
    selects = []
    for _, (table_name, _) in TABLES.items():
        selects.append(f"  SELECT '{table_name}' tabla, COUNT(*) filas FROM {table_name}")
    lines.append("\n  UNION ALL\n".join(selects))
    lines.extend((" ) ORDER BY tabla;", "", f"-- Total esperado en las 19 tablas: {total}", ""))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Generado {OUTPUT}: {total} INSERTs")


if __name__ == "__main__":
    main()
